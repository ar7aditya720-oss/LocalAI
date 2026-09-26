import os
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.config import WORKSPACE_DIR

class DocumentTool:
    @staticmethod
    def create_approval_note(
        title: str,
        subject: str,
        reference_no: str,
        department: str,
        key_findings: List[str],
        recommendations: List[str],
        output_filename: str = "Approval_Note.docx"
    ) -> Dict[str, Any]:
        """
        Generates an official Word Document Approval Note (.docx).
        """
        doc = docx.Document()
        
        # Heading / Title
        header_p = doc.add_paragraph()
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = header_p.add_run("OFFICIAL APPROVAL NOTE & TECHNICAL BRIEF")
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x0D, 0x11, 0x17)
        
        doc.add_paragraph() # Spacer
        
        # Metadata Table
        table = doc.add_table(rows=4, cols=2)
        table.style = 'Table Grid'
        
        meta = [
            ("Document Reference:", reference_no),
            ("Subject:", subject),
            ("Department / Unit:", department),
            ("Date & Time:", "2026-09-23 | Sovereign LocalAI System")
        ]
        
        for i, (label, val) in enumerate(meta):
            row = table.rows[i]
            row.cells[0].paragraphs[0].add_run(label).bold = True
            row.cells[1].paragraphs[0].add_run(val)
            
        doc.add_paragraph()
        
        # Section 1: Executive Summary
        h1 = doc.add_heading("1. Executive Summary", level=1)
        h1.runs[0].font.color.rgb = RGBColor(0x2B, 0x54, 0x9A)
        
        p = doc.add_paragraph(
            f"This approval note has been automatically compiled by LocalAI Sovereign Workbench following "
            f"on-premise document analysis and verification against organizational SOPs for {subject}."
        )
        
        # Section 2: Key Technical Findings
        h2 = doc.add_heading("2. Key Technical Findings", level=1)
        h2.runs[0].font.color.rgb = RGBColor(0x2B, 0x54, 0x9A)
        
        for finding in key_findings:
            doc.add_paragraph(finding, style='List Bullet')
            
        # Section 3: Recommendations & Action Plan
        h3 = doc.add_heading("3. Recommendations & Next Actions", level=1)
        h3.runs[0].font.color.rgb = RGBColor(0x2B, 0x54, 0x9A)
        
        for rec in recommendations:
            doc.add_paragraph(rec, style='List Bullet')
            
        # Sign-off section
        doc.add_paragraph()
        p_sign = doc.add_paragraph()
        p_sign.add_run("Prepared & Verified by: ").bold = True
        p_sign.add_run("LocalAI Sovereign Agent Engine\n")
        p_sign.add_run("Approval Status: ").bold = True
        p_sign.add_run("RECOMMENDED FOR CHIEF ENGINEER / AUTHORITY REVIEW")
        
        # Save file
        out_path = os.path.join(str(WORKSPACE_DIR), output_filename)
        doc.save(out_path)
        
        return {
            "status": "SUCCESS",
            "filename": output_filename,
            "filepath": out_path,
            "file_size_bytes": os.path.getsize(out_path)
        }
