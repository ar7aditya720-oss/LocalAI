import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os
from typing import Dict, Any, List
from app.config import WORKSPACE_DIR

class SpreadsheetTool:
    @staticmethod
    def create_inspection_sheet(
        title: str,
        rows_data: List[List[Any]],
        headers: List[str] = None,
        output_filename: str = "Inspection_Report_Matrix.xlsx"
    ) -> Dict[str, Any]:
        """
        Creates a formatted Excel spreadsheet (.xlsx).
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LocalAI Data Sheet"
        
        # Default Headers if not provided
        if not headers:
            headers = ["Item ID", "Equipment Tag", "Finding / Defect", "Severity", "SOP Reference", "Action Required"]
            
        # Title Header Row
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = title.upper()
        title_cell.font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
        title_cell.fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 30
        
        # Column Headers
        ws.append([]) # Empty row 2
        ws.append(headers) # Row 3
        
        header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        header_font = Font(name='Calibri', size=11, bold=True, color='1F4E79')
        
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
        ws.row_dimensions[3].height = 24
        
        # Append rows
        for row in rows_data:
            ws.append(row)
            
        # Format columns width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
            
        out_path = os.path.join(str(WORKSPACE_DIR), output_filename)
        wb.save(out_path)
        
        return {
            "status": "SUCCESS",
            "filename": output_filename,
            "filepath": out_path,
            "file_size_bytes": os.path.getsize(out_path)
        }
